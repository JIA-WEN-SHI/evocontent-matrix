from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Octopus Excel exports to KB payload JSON.",
    )
    parser.add_argument("--input-dir", required=True, help="Directory containing .xlsx files")
    parser.add_argument("--output-file", required=True, help="Output payload JSON path")
    parser.add_argument("--entity-type", default="case", choices=["case", "asset", "user_need", "review"])
    parser.add_argument("--domain-slug", default="japan_immigration")
    parser.add_argument("--source", default="octopus")
    parser.add_argument("--source-run-id", default="")
    parser.add_argument("--max-items", type=int, default=500)
    parser.add_argument("--disable-tags", action="store_true", help="Do not emit tags to avoid tag upsert conflicts")
    return parser.parse_args()


def as_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def to_json_safe(v: Any) -> Any:
    if isinstance(v, datetime):
        if v.tzinfo is None:
            return v.replace(tzinfo=timezone.utc).astimezone().isoformat()
        return v.astimezone().isoformat()
    if isinstance(v, (list, tuple)):
        return [to_json_safe(x) for x in v]
    if isinstance(v, dict):
        return {str(k): to_json_safe(val) for k, val in v.items()}
    return v


def as_int(v: Any) -> int | None:
    s = as_text(v)
    if not s:
        return None
    s = s.replace(",", "")
    m = re.search(r"-?\d+", s)
    if not m:
        return None
    try:
        return int(m.group(0))
    except ValueError:
        return None


def extract_note_id(url: str) -> str:
    if not url:
        return ""
    m = re.search(r"/explore/([0-9a-zA-Z]+)", url)
    if m:
        return m.group(1)
    return ""


def extract_tags(text: str) -> list[str]:
    if not text:
        return []
    tags = re.findall(r"#([^\s#]+)", text)
    out: list[str] = []
    seen: set[str] = set()
    for t in tags:
        t2 = t.strip().strip("#")
        if not t2:
            continue
        # Normalize ASCII tags to lowercase to reduce case-sensitive collisions
        t_norm = t2.lower() if re.search(r"[A-Za-z]", t2) else t2
        key = t_norm.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(t_norm)
    return out


def contains_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def load_rows_from_excel(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    try:
        headers = [as_text(x) for x in next(iterator)]
    except StopIteration:
        return []
    rows: list[dict[str, Any]] = []
    for row in iterator:
        data = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        rows.append(data)
    return rows


def map_row_to_item(row: dict[str, Any], disable_tags: bool = False) -> dict[str, Any] | None:
    title = as_text(row.get("标题"))
    content = as_text(row.get("正文内容"))
    url = as_text(row.get("笔记链接")) or as_text(row.get("发布地址")) or as_text(row.get("视频链接/首图链接"))
    author = as_text(row.get("博主名称"))
    post_type = as_text(row.get("笔记类型"))
    publish_time = as_text(row.get("发布时间"))
    keyword = as_text(row.get("关键词"))
    likes = as_int(row.get("点赞数"))
    collects = as_int(row.get("收藏数"))
    comments = as_int(row.get("评论数"))

    # Minimal quality guard to avoid empty rows from Excel
    if not (title or content or url):
        return None
    if not contains_chinese(f"{title} {content}"):
        return None

    note_id = extract_note_id(url)
    source_ref = f"xhs:{note_id}" if note_id else (f"url:{url}" if url else "")
    tags = []
    if not disable_tags:
        tags = extract_tags(content)
        if keyword:
            tags = [keyword] + [t for t in tags if t != keyword]

    metrics = {
        "likes": likes,
        "collects": collects,
        "comments_count": comments,
        "shares": None,
        "impressions": None,
    }

    now_iso = datetime.now(timezone.utc).astimezone().isoformat()
    item = {
        "title": title,
        "content": content or title,
        "url": url,
        "author": author,
        "platform": "xiaohongshu",
        "metrics": metrics,
        "tags": tags,
        "captured_at": now_iso,
        "raw": {
            **to_json_safe(row),
            "note_id": note_id,
            "publish_time": publish_time,
            "post_type": post_type,
            "keyword": keyword,
        },
        "source_ref": source_ref or f"row:{title[:20]}",
    }
    return item


def main() -> int:
    args = parse_args()
    input_dir = Path(args.input_dir)
    output_file = Path(args.output_file)
    if not input_dir.exists():
        raise SystemExit(f"input dir not found: {input_dir}")

    files = sorted([p for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$")], key=lambda x: x.name)
    if not files:
        raise SystemExit(f"no .xlsx files found in {input_dir}")

    mapped: list[dict[str, Any]] = []
    for f in files:
        rows = load_rows_from_excel(f)
        for r in rows:
            item = map_row_to_item(r, disable_tags=args.disable_tags)
            if item:
                mapped.append(item)

    # Deduplicate by source_ref, keep the one with higher interaction sum
    best: dict[str, dict[str, Any]] = {}
    for item in mapped:
        key = item["source_ref"]
        score = sum(
            x or 0
            for x in [
                item["metrics"].get("likes"),
                item["metrics"].get("collects"),
                item["metrics"].get("comments_count"),
                item["metrics"].get("shares"),
            ]
        )
        existed = best.get(key)
        if not existed:
            best[key] = item
            continue
        existed_score = sum(
            x or 0
            for x in [
                existed["metrics"].get("likes"),
                existed["metrics"].get("collects"),
                existed["metrics"].get("comments_count"),
                existed["metrics"].get("shares"),
            ]
        )
        if score >= existed_score:
            best[key] = item

    items = list(best.values())
    if len(items) > args.max_items:
        items = items[: args.max_items]

    source_run_id = args.source_run_id or f"octopus-excel-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
    payload = {
        "domain_slug": args.domain_slug,
        "account_id": "",
        "source": args.source,
        "source_run_id": source_run_id,
        "entity_type": args.entity_type,
        "items": items,
    }

    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"excel_files={len(files)}")
    print(f"rows_mapped={len(mapped)}")
    print(f"rows_deduped={len(items)}")
    print(f"output={output_file}")
    print(f"source_run_id={source_run_id}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
